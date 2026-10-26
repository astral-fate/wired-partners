#!/usr/bin/env python3
"""
Reusable toolkit for building a presentation-grade technical demo deck with
python-pptx, in the navy + teal design system.

Import everything:  from deck_helpers import *

Every primitive takes `prs` or `slide` explicitly, so a build script stays a
flat, readable sequence of calls (see example_security_deck.py for a full deck).

Key design decisions baked in here:
  * 20 x 11.25 in stock (16:9 at a large scale — matches a real pitch deck and
    gives room for big display numbers).
  * A literal "\\n" in run text does NOT center per line in LibreOffice/PowerPoint.
    `text()` therefore emits a real <a:br/> element between segments, so paragraph
    alignment applies to every line. This is the single most common render bug.
  * QR codes are dark-navy on white (high contrast, on-brand, phone-scannable).
"""
from __future__ import annotations
import pathlib
import qrcode
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---- palette (adapted from a professional navy/teal pitch deck) ------------
NAVY   = RGBColor(0x1B, 0x2A, 0x3A)   # primary dark ground / dark text
NAVY2  = RGBColor(0x0B, 0x25, 0x45)   # deeper navy (table last row, etc.)
INK    = RGBColor(0x0F, 0x17, 0x2A)   # near-black body text on light
CARDDK = RGBColor(0x22, 0x33, 0x45)   # dark card on navy
TEAL   = RGBColor(0x0F, 0x76, 0x6E)   # deep teal (kickers, labels)
TEALBR = RGBColor(0x14, 0xB8, 0xA6)   # bright teal accent (rules, highlights)
TEALLT = RGBColor(0xCC, 0xFB, 0xF1)   # pale teal fill
MIST   = RGBColor(0xCF, 0xE3, 0xE8)   # light text on navy
SLATE  = RGBColor(0x47, 0x55, 0x69)   # secondary body on light
SLATE2 = RGBColor(0x5F, 0x6F, 0x80)   # tertiary / footer
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
RED    = RGBColor(0xDC, 0x26, 0x26)   # FAIL / high attack success
AMBER  = RGBColor(0xF5, 0x9E, 0x0B)   # ATTENTION / caution
GREEN  = RGBColor(0x15, 0x80, 0x3D)   # PASS / resisted
PAPER  = RGBColor(0xF4, 0xF7, 0xF9)   # light slide ground
CARD   = RGBColor(0xFF, 0xFF, 0xFF)
LINEC  = RGBColor(0xD8, 0xE1, 0xE7)   # hairline borders
# tints for dark cards
POS_ON_DARK = RGBColor(0x5E, 0xE0, 0xB0)   # green text on navy
NEG_ON_DARK = RGBColor(0xFF, 0x9C, 0x9C)   # red text on navy

# fonts (Segoe UI ships on Windows; Consolas for code/monospace)
HEAD  = "Segoe UI Semibold"   # kickers / eyebrow labels
HEADB = "Segoe UI"            # big bold titles & display numbers (set bold=True)
BODY  = "Segoe UI"
MONO  = "Consolas"

SW, SH = 20.0, 11.25          # slide size in inches


def new_deck() -> Presentation:
    prs = Presentation()
    prs.slide_width  = Inches(SW)
    prs.slide_height = Inches(SH)
    return prs


# ---- QR ---------------------------------------------------------------------
def make_qr(url: str, out_path, dark: str = "#1B2A3A") -> pathlib.Path:
    """Generate a scannable QR PNG (dark navy on white)."""
    q = qrcode.QRCode(border=2, box_size=12,
                      error_correction=qrcode.constants.ERROR_CORRECT_M)
    q.add_data(url); q.make(fit=True)
    img = q.make_image(fill_color=dark, back_color="white").convert("RGB")
    p = pathlib.Path(out_path); p.parent.mkdir(parents=True, exist_ok=True)
    img.save(p)
    return p


# ---- primitives -------------------------------------------------------------
def slide(prs, bg=PAPER):
    s = prs.slides.add_slide(prs.slide_layouts[6])   # blank layout
    s.background.fill.solid(); s.background.fill.fore_color.rgb = bg
    return s


def _set_line_spacing(para, mult):
    pPr = para._p.get_or_add_pPr()
    e = pPr.find(qn("a:lnSpc"))
    if e is not None: pPr.remove(e)
    ln = pPr.makeelement(qn("a:lnSpc"), {})
    ln.append(ln.makeelement(qn("a:spcPct"), {"val": str(int(mult * 100000))}))
    pPr.insert(0, ln)


def text(s, l, t, w, h, runs, size=18, color=INK, bold=False, font=BODY,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.0, wrap=True,
         space_after=0.0):
    """Add a textbox.

    `runs` accepts three shapes, richest last:
      * "plain string"                              -> one paragraph
      * [("bold bit", {"bold":True}), ("rest", {})] -> one paragraph, many runs
      * [[("line 1 run", {})], [("line 2 run", {})]]-> many paragraphs
    A "\\n" inside any run text becomes a real <a:br/> so alignment holds per line.
    Per-run overrides in the dict: size, bold, font, color, italic.
    """
    tb = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = wrap; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    if isinstance(runs, str):
        paras = [[(runs, {})]]
    elif runs and isinstance(runs[0], tuple):
        paras = [runs]
    else:
        paras = runs
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if spacing != 1.0: _set_line_spacing(p, spacing)
        p.space_after = Pt(space_after) if space_after else Pt(0)
        if isinstance(para, str): para = [(para, {})]
        for txt, o in para:
            for si, seg in enumerate(str(txt).split("\n")):
                if si > 0:
                    p._p.append(p._p.makeelement(qn("a:br"), {}))
                r = p.add_run(); r.text = seg
                r.font.size = Pt(o.get("size", size))
                r.font.bold = o.get("bold", bold)
                r.font.name = o.get("font", font)
                r.font.color.rgb = o.get("color", color)
                if o.get("italic"): r.font.italic = True
    return tb


def rect(s, l, t, w, h, fill=CARD, line=None, line_w=1.0, radius=0.10,
         shadow=False, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    sp = s.shapes.add_shape(shape, Inches(l), Inches(t), Inches(w), Inches(h))
    if fill is None: sp.fill.background()
    else: sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None: sp.line.fill.background()
    else: sp.line.color.rgb = line; sp.line.width = Pt(line_w)
    sp.shadow.inherit = False
    if shadow:
        el = sp._element.spPr
        ef = el.makeelement(qn("a:effectLst"), {})
        sh = ef.makeelement(qn("a:outerShdw"),
                            {"blurRad": "90000", "dist": "40000",
                             "dir": "5400000", "rotWithShape": "0"})
        clr = sh.makeelement(qn("a:srgbClr"), {"val": "1B2A3A"})
        clr.append(clr.makeelement(qn("a:alpha"), {"val": "18000"}))
        sh.append(clr); ef.append(sh); el.append(ef)
    try: sp.adjustments[0] = radius
    except Exception: pass
    return sp


def line(s, l, t, w, h, color=TEALBR):
    """A thin filled bar — use for accent rules and dividers."""
    sp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(t),
                            Inches(max(w, 0.001)), Inches(max(h, 0.001)))
    sp.fill.solid(); sp.fill.fore_color.rgb = color
    sp.line.fill.background(); sp.shadow.inherit = False
    return sp


def kicker(s, txt, l=0.9, t=0.62, color=TEAL):
    """Eyebrow label above a slide title."""
    text(s, l, t, 16, 0.5, txt.upper(), size=15, color=color, bold=True, font=HEAD)


def title(s, txt, l=0.9, t=1.02, w=18.2, color=INK, size=40):
    text(s, l, t, w, 1.4, txt, size=size, color=color, bold=True, font=HEADB)


def footer(s, n, dark=False, brand="Data Intelligence SDK  \u00b7  Demo"):
    c = MIST if dark else SLATE2
    text(s, 0.9, SH - 0.62, 12, 0.4, brand, size=11, color=c, font=BODY)
    text(s, SW - 3.9, SH - 0.62, 3.0, 0.4, f"{n:02d}", size=11, color=c,
         font=BODY, align=PP_ALIGN.RIGHT)


def chip(s, l, t, txt, fill, tcolor=WHITE, w=None, size=12):
    w = w or (0.28 + 0.105 * len(txt))
    rect(s, l, t, w, 0.42, fill=fill, radius=0.5)
    text(s, l, t + 0.02, w, 0.38, txt, size=size, color=tcolor, bold=True,
         font=HEAD, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    return w


def qr_block(s, qr_png, l, t, size, caption, cap_color=SLATE, cap_size=11.5):
    """Place a QR image with a centered caption beneath it."""
    s.shapes.add_picture(str(qr_png), Inches(l), Inches(t), Inches(size), Inches(size))
    text(s, l - 0.4, t + size + 0.06, size + 0.8, 0.6, caption, size=cap_size,
         color=cap_color, align=PP_ALIGN.CENTER, font=BODY, spacing=0.95)


def divider(prs, num, kick, big):
    """Full-bleed navy section divider: giant ghost number + kicker + title.
    `big` may contain "\\n" for a two-line title (renders correctly via <a:br/>)."""
    s = slide(prs, NAVY)
    line(s, 0, 0, SW, 0.22, TEALBR)
    text(s, 1.2, 4.0, 3.0, 2.0, num, size=150, color=CARDDK, bold=True, font=HEADB)
    text(s, 1.3, 3.6, 16, 0.6, kick.upper(), size=18, color=TEALBR, bold=True, font=HEAD)
    text(s, 1.25, 4.25, 17, 2.4, big, size=52, color=WHITE, bold=True, font=HEADB, spacing=1.02)
    return s


# ---- render-verify (MANDATORY before delivery) ------------------------------
def render_to_pngs(pptx_path, out_dir, dpi=110):
    """LibreOffice headless -> PDF -> per-page PNG. Returns list of PNG paths.

    PowerPoint COM automation is blocked in this environment; LibreOffice
    headless is the reliable renderer. Install once:
        winget install TheDocumentFoundation.LibreOffice
    Verify decks by READING the PNGs — never by reading the XML. Layout bugs
    (overflow, collisions, mis-centered lines) only show in the render.
    """
    import subprocess, pymupdf
    out = pathlib.Path(out_dir); out.mkdir(parents=True, exist_ok=True)
    soffice = r"C:\Program Files\LibreOffice\program\soffice.exe"
    subprocess.run([soffice, "--headless", "--convert-to", "pdf",
                    "--outdir", str(out), str(pptx_path)],
                   check=True, capture_output=True, timeout=180)
    pdf = out / (pathlib.Path(pptx_path).stem + ".pdf")
    doc = pymupdf.open(str(pdf))
    pngs = []
    for i, pg in enumerate(doc, 1):
        p = out / f"slide_{i:02d}.png"
        pg.get_pixmap(dpi=dpi).save(str(p)); pngs.append(p)
    return pngs


def decode_qr(png_path):
    """Decode-verify a QR PNG with OpenCV (upscales, since tiny QRs falsely fail).
    Returns the decoded string or None."""
    import cv2
    det = cv2.QRCodeDetector()
    for fx in (3, 5, 8, 10):
        for interp in (cv2.INTER_CUBIC, cv2.INTER_LANCZOS4, cv2.INTER_NEAREST):
            img = cv2.imread(str(png_path))
            img = cv2.resize(img, None, fx=fx, fy=fx, interpolation=interp)
            data, _, _ = det.detectAndDecode(img)
            if data: return data
    return None
